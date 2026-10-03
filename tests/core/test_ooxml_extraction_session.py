"""Core OOXML output parity, streaming and resource lifetime (#149)."""

from datetime import datetime, time, timedelta
from pathlib import Path
from unittest.mock import Mock
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from defusedxml.common import EntitiesForbidden
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.datetime import MAC_EPOCH
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table
import pytest

from exstruct.core import ooxml_session
from exstruct.core.backends.ooxml_backend import OoxmlRichBackend
from exstruct.core.backends.openpyxl_backend import OpenpyxlBackend
from exstruct.core.ooxml_session import OoxmlExtractionSession

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def make_book(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "O'Brien, data"
    ws.append(["Name", "Count"])
    ws.append(["item", 7])
    ws["A2"].hyperlink = "https://example.com/a?q=1#top"
    for coord, value in {
        "D1": "merged",
        "D3": True,
        "E3": False,
        "F3": "0012",
        "G3": "1.250",
        "H3": " NA ",
        "I3": "NA",
        "J3": "#DIV/0!",
        "A5": datetime(2026, 1, 2, 3, 4, 5),
        "B5": time(12, 30),
        "C5": timedelta(days=2, hours=3),
        "D5": -1.5,
        "A8": "=SUM(B2)",
        "B8": "=A2",
        "A10000": "sparse",
    }.items():
        ws[coord] = value
    ws["I3"].hyperlink = "https://example.com/filtered"
    ws["A7"] = "NA"
    ws["A7"].hyperlink = "https://example.com/only-filtered"
    ws["B2"].hyperlink = Hyperlink(ref="B2", location="'Other'!A1")
    ws["C8"] = ArrayFormula(ref="C8:C9", text="=SUM(B2)")
    ws.merge_cells("D1:E2")
    ws.merge_cells("D10:E10")
    ws.print_area = ["A1:B5", "D1:E2"]
    ws.add_table(Table(displayName="Items", ref="A1:B2"))
    chart = BarChart()
    chart.add_data(
        Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True
    )
    ws.add_chart(chart, "L1")
    wb.defined_names.add(DefinedName("Global", attr_text="'O''Brien, data'!$A$1"))
    other = wb.create_sheet("Other")
    other.sheet_state = "hidden"
    wb.save(path)
    wb.close()


def replace_parts(path: Path, parts: dict[str, bytes]) -> None:
    with ZipFile(path) as archive:
        payload = {name: archive.read(name) for name in archive.namelist()}
    payload.update(parts)
    with ZipFile(path, "w") as archive:
        for name, data in payload.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
@pytest.mark.parametrize("include_links", [False, True])
def test_core_parity(tmp_path: Path, suffix: str, include_links: bool) -> None:
    path = tmp_path / f"book{suffix}"
    make_book(path)
    expected = OpenpyxlBackend(path)
    with OoxmlExtractionSession(path) as session:
        assert session.extract_cells(
            include_links=include_links
        ) == expected.extract_cells(include_links=include_links)
        assert session.extract_formulas_map() == expected.extract_formulas_map()
        assert session.extract_print_areas() == expected.extract_print_areas()
        actual_merged = session.extract_merged_cells()
        expected_merged = expected.extract_merged_cells()
        for name in actual_merged:
            assert set(actual_merged[name]) == set(expected_merged[name])
        tables = session.extract_explicit_tables()
        assert [t.ref for t in tables["O'Brien, data"]] == ["A1:B2"]
        assert tables["O'Brien, data"][0].columns == ("Name", "Count")
        assert tables["O'Brien, data"][0].display_name == "Items"
        assert tables["Other"] == []
        assert [(s.name, s.state) for s in session.sheets()] == [
            ("O'Brien, data", "visible"),
            ("Other", "hidden"),
        ]
        assert any(
            n.name == "Global" and n.local_sheet_id is None
            for n in session.defined_names()
        )


def test_shared_strings_cached_formulas_and_shared_formula_parity(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    with ZipFile(path) as archive:
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
    ET.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        PartName="/xl/sharedStrings.xml",
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    )
    ET.SubElement(
        rels,
        f"{{{PKG}}}Relationship",
        Id="strings",
        Type=f"{REL}/sharedStrings",
        Target="sharedStrings.xml",
    )
    sheet = f'''<worksheet xmlns="{MAIN}"><sheetData>
      <row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="inlineStr"><is><r><t>in</t></r><r><t>line</t></r><rPh sb="0" eb="1"><t>phonetic</t></rPh></is></c></row>
      <row r="2"><c r="A2"><f t="shared" si="0" ref="A2:A3">B2+$C$1</f><v>4</v></c><c r="B2" t="str"><f>TEXT(1,"0")</f><v>cached text</v></c></row>
      <row r="3"><c r="A3"><f t="shared" si="0"/><v>5</v></c><c r="B3" t="b"><v>1</v></c></row>
      <row r="4"><c r="A4" t="e"><v>#N/A</v></c><c r="B4" t="d"><v>2026-01-02T03:04:05</v></c></row>
      <row r="5"><c r="A5"><f t="array" ref="A5:A6">SUM(B2)</f><v>2</v></c></row>
    </sheetData><mergeCells><mergeCell ref="A1:C1"/></mergeCells></worksheet>'''
    # Keep the rich inline string outside the merge extent.
    sheet = sheet.replace('ref="A1:C1"', 'ref="D1:E1"')
    strings = f'''<sst xmlns="{MAIN}" count="1" uniqueCount="1"><si><r><t xml:space="preserve">shared </t></r><r><t>string</t></r><rPh sb="0" eb="1"><t>ignored</t></rPh></si></sst>'''
    replace_parts(
        path,
        {
            "xl/_rels/workbook.xml.rels": ET.tostring(rels),
            "xl/worksheets/sheet1.xml": sheet.encode(),
            "xl/sharedStrings.xml": strings.encode(),
            "[Content_Types].xml": ET.tostring(content_types),
        },
    )
    expected = OpenpyxlBackend(path)
    with OoxmlExtractionSession(path) as session:
        assert session.extract_cells() == expected.extract_cells(include_links=False)
        assert session.extract_formulas_map() == expected.extract_formulas_map()
        rows = session.extract_cells()["O'Brien, data"]
        assert rows[0].c == {"0": "shared string", "1": "inline"}
        assert session.extract_formulas_map().sheets["O'Brien, data"].formulas_map[
            "=B3+$C$1"
        ] == [(3, 0)]


def test_archive_and_worksheet_reads_reused_and_rich_parity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    standalone = OoxmlRichBackend(path)
    expected_charts = standalone.extract_charts(mode="light")
    opened: list[ZipFile] = []
    real_zip = ZipFile

    def track(path: Path) -> ZipFile:
        archive = real_zip(path)
        opened.append(archive)
        return archive

    monkeypatch.setattr(ooxml_session, "ZipFile", track)
    with OoxmlExtractionSession(path) as session:
        assert opened == []
        session.extract_cells()
        archive = session.archive()
        original_open = archive.open
        reads: list[str] = []

        def read(name: str, *args: object, **kwargs: object) -> object:
            reads.append(name)
            return original_open(name, *args, **kwargs)

        monkeypatch.setattr(archive, "open", read)
        session.extract_cells(include_links=True)
        session.extract_formulas_map()
        session.extract_merged_cells()
        session.extract_print_areas()
        session.extract_explicit_tables()
        assert "xl/worksheets/sheet1.xml" not in reads
        rich = OoxmlRichBackend(path, session=session)
        assert rich.extract_charts(mode="light") == expected_charts
        assert rich.extract_shapes(mode="light") == standalone.extract_shapes(
            mode="light"
        )
        assert session.read_drawings() is session.read_drawings()
        assert len(opened) == 1
    assert opened[0].fp is None


def test_exception_closes_and_reuse_rejected(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    replace_parts(path, {"xl/worksheets/sheet1.xml": b"<broken>"})
    session = OoxmlExtractionSession(path)
    with pytest.raises(ET.ParseError), session:
        archive = session.archive()
        session.extract_cells()
    assert archive.fp is None
    session.close()
    for call in (
        session.sheets,
        session.extract_cells,
        session.archive,
        session.__enter__,
    ):
        with pytest.raises(RuntimeError, match="closed"):
            call()


def test_defused_stream_rejects_entities(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    replace_parts(
        path,
        {
            "xl/worksheets/sheet1.xml": b'<!DOCTYPE x [<!ENTITY e "expanded">]><x>&e;</x>'
        },
    )
    with pytest.raises(EntitiesForbidden), OoxmlExtractionSession(path) as session:
        session.extract_cells()


def test_worksheet_stream_never_uses_archive_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    with OoxmlExtractionSession(path) as session:
        archive = session.archive()
        original = archive.read

        def read(name: str, *args: object, **kwargs: object) -> bytes:
            assert not ("/worksheets/" in name and name.endswith(".xml"))
            return original(name, *args, **kwargs)

        monkeypatch.setattr(archive, "read", read)
        monkeypatch.setattr(
            "openpyxl.load_workbook",
            Mock(side_effect=AssertionError("no workbook model")),
        )
        assert session.extract_cells()["O'Brien, data"][-1].r == 10000


def test_sessions_observe_updates_and_reject_biff(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    make_book(path)
    with OoxmlExtractionSession(path) as first:
        previous = first.extract_cells()
    make_book(path)
    replace_parts(
        path,
        {
            "xl/worksheets/sheet2.xml": f'<worksheet xmlns="{MAIN}"><sheetData><row r="2"><c r="A2" t="inlineStr"><is><t>new</t></is></c></row></sheetData></worksheet>'.encode()
        },
    )
    with OoxmlExtractionSession(path) as second:
        assert previous != second.extract_cells()
    with OoxmlExtractionSession(tmp_path / "book.xls") as session:
        with pytest.raises(ValueError, match="xlsx"):
            session.extract_cells()


def test_implicit_coordinates_and_out_of_range_date_parity(tmp_path: Path) -> None:
    path = tmp_path / "implicit.xlsx"
    make_book(path)
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    style = root.find(f"{{{MAIN}}}sheetData/{{{MAIN}}}row[@r='5']/{{{MAIN}}}c[@r='A5']")
    assert style is not None
    date_style = style.attrib["s"]
    replace_parts(
        path,
        {
            "xl/worksheets/sheet1.xml": f'''<worksheet xmlns="{MAIN}"><sheetData>
      <row><c t="inlineStr"><is><t>first</t></is></c><c r="D1"><v>2</v></c><c><f>D1+1</f><v>3</v></c></row>
      <row r="4"><c><v>4</v></c><c s="{date_style}"><v>100000000</v></c></row>
      <row><c><v>5</v></c></row>
    </sheetData></worksheet>'''.encode()
        },
    )
    with OoxmlExtractionSession(path) as session:
        backend = OpenpyxlBackend(path)
        with pytest.warns(UserWarning, match="outside the limits"):
            assert session.extract_cells() == backend.extract_cells(include_links=False)
        with pytest.warns(UserWarning, match="outside the limits"):
            assert session.extract_formulas_map() == backend.extract_formulas_map()
        assert session.extract_cells()["O'Brien, data"][0].c == {
            "0": "first",
            "3": 2,
            "4": 3,
        }


def test_mac_epoch_and_iso_date_parity(tmp_path: Path) -> None:
    path = tmp_path / "epoch.xlsx"
    wb = Workbook()
    wb.epoch = MAC_EPOCH
    ws = wb.active
    assert ws is not None
    ws["A1"] = datetime(2026, 1, 2)
    ws["B1"] = time(6, 0)
    wb.save(path)
    wb.close()
    with OoxmlExtractionSession(path) as session:
        assert session.extract_cells() == OpenpyxlBackend(path).extract_cells(
            include_links=False
        )


def test_foreign_rich_session_rejected_and_unresolved_formula_skipped(
    tmp_path: Path,
) -> None:
    path = tmp_path / "formulas.xlsx"
    make_book(path)
    replace_parts(
        path,
        {
            "xl/worksheets/sheet1.xml": f'''<worksheet xmlns="{MAIN}"><sheetData>
    <row r="1"><c r="A1"><f t="shared" si="missing"/><v>1</v></c></row>
    </sheetData></worksheet>'''.encode()
        },
    )
    with OoxmlExtractionSession(path) as session:
        assert session.extract_formulas_map().sheets["O'Brien, data"].formulas_map == {}
        assert session.extract_cells()["O'Brien, data"][0].c == {"0": 1}
        with pytest.raises(ValueError, match="different workbook"):
            OoxmlRichBackend(tmp_path / "foreign.xlsx", session=session)
