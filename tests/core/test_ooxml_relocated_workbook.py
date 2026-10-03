"""Regression coverage for OOXML workbooks outside the conventional xl path."""

from pathlib import Path
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
import pytest

from exstruct.core import ooxml_session
from exstruct.core.backends.ooxml_backend import OoxmlRichBackend
from exstruct.core.ooxml_session import OoxmlExtractionSession

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWINGML = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _make_chart_workbook(path: Path) -> None:
    """Build identical chart-bearing workbooks for layout comparisons."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Category", "Value"])
    sheet.append(["A", 7])
    sheet.append(["B", 12])

    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D1")
    workbook.save(path)
    workbook.close()


def _append_shape(drawing_xml: bytes) -> bytes:
    """Add a deterministic shape anchor beside the generated chart."""
    ET.register_namespace("xdr", DRAWING)
    ET.register_namespace("a", DRAWINGML)
    root = ET.fromstring(drawing_xml)
    root.append(
        ET.fromstring(
            f"""
            <xdr:twoCellAnchor xmlns:xdr="{DRAWING}" xmlns:a="{DRAWINGML}">
              <xdr:from>
                <xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff>
                <xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff>
              </xdr:from>
              <xdr:to>
                <xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff>
                <xdr:row>6</xdr:row><xdr:rowOff>0</xdr:rowOff>
              </xdr:to>
              <xdr:sp>
                <xdr:nvSpPr>
                  <xdr:cNvPr id="99" name="Regression Shape" />
                  <xdr:cNvSpPr />
                </xdr:nvSpPr>
                <xdr:spPr>
                  <a:prstGeom prst="rect"><a:avLst /></a:prstGeom>
                </xdr:spPr>
                <xdr:txBody>
                  <a:bodyPr /><a:lstStyle />
                  <a:p><a:r><a:rPr lang="en-US"/><a:t>Relocated shape</a:t></a:r></a:p>
                </xdr:txBody>
              </xdr:sp>
              <xdr:clientData />
            </xdr:twoCellAnchor>
            """
        )
    )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _add_shape_to_workbook(path: Path) -> None:
    """Inject the same shape into each workbook before relocating its parts."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/drawings/drawing1.xml"] = _append_shape(parts["xl/drawings/drawing1.xml"])
    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _relocate_workbook_part(path: Path) -> None:
    """Move the workbook and repair root, part and content-type references."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    parts["custom/book.xml"] = parts.pop("xl/workbook.xml")

    workbook_rels = ET.fromstring(parts.pop("xl/_rels/workbook.xml.rels"))
    for relationship in workbook_rels.findall(f"{{{PKG}}}Relationship"):
        if relationship.get("TargetMode") == "External":
            continue
        target = relationship.get("Target")
        if not target:
            continue
        old_part_path = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        relationship.set("Target", posixpath.relpath(old_part_path, "custom"))
    parts["custom/_rels/book.xml.rels"] = ET.tostring(workbook_rels)

    root_rels = ET.fromstring(parts["_rels/.rels"])
    for relationship in root_rels.findall(f"{{{PKG}}}Relationship"):
        if relationship.get("Type") == f"{REL}/officeDocument":
            relationship.set("Target", "custom/book.xml")
            break
    parts["_rels/.rels"] = ET.tostring(root_rels)

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for override in content_types.findall(f"{{{CONTENT_TYPES}}}Override"):
        if override.get("PartName") == "/xl/workbook.xml":
            override.set("PartName", "/custom/book.xml")
            break
    parts["[Content_Types].xml"] = ET.tostring(content_types)

    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _read_outputs(path: Path) -> tuple[object, object, object, object]:
    """Exercise drawings first, then core and rich consumers of the same ZIP."""
    with OoxmlExtractionSession(path) as session:
        drawings = session.read_drawings()
        cells = session.extract_cells()
        rich = OoxmlRichBackend(path, session=session)
        charts = rich.extract_charts(mode="light")
        shapes = rich.extract_shapes(mode="light")
    return cells, charts, shapes, drawings


def test_relocated_workbook_drawings_first_matches_standard_layout(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Retain cells, charts and shapes with one archive per relocated session."""
    standard_path = tmp_path / "standard.xlsx"
    relocated_path = tmp_path / "relocated.xlsx"
    _make_chart_workbook(standard_path)
    _make_chart_workbook(relocated_path)
    _add_shape_to_workbook(standard_path)
    _add_shape_to_workbook(relocated_path)
    _relocate_workbook_part(relocated_path)

    opened: list[ZipFile] = []
    real_zip = ZipFile

    def track(path: Path) -> ZipFile:
        archive = real_zip(path)
        opened.append(archive)
        return archive

    monkeypatch.setattr(ooxml_session, "ZipFile", track)
    standard_outputs = _read_outputs(standard_path)
    relocated_outputs = _read_outputs(relocated_path)

    assert relocated_outputs == standard_outputs
    assert standard_outputs[1]["Data"]
    assert standard_outputs[2]["Data"]
    assert standard_outputs[3]["Data"].charts
    assert standard_outputs[3]["Data"].shapes
    assert len(opened) == 2
    assert all(archive.fp is None for archive in opened)
