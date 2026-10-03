"""Extraction resource sharing, isolation and compatibility contracts (#146)."""

from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from dataclasses import replace
from pathlib import Path
from threading import Barrier
from types import ModuleType, SimpleNamespace
from typing import Any
from unittest.mock import Mock

from openpyxl import Workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.worksheet.table import Table
import pytest

from exstruct.core import cells, pipeline, workbook
from exstruct.core.backends import openpyxl_backend
from exstruct.core.backends.openpyxl_backend import OpenpyxlBackend
from exstruct.core.integrate import extract_workbook
from exstruct.core.openpyxl_session import OpenpyxlExtractionSession
from exstruct.core.pipeline import ExtractionInputs, ExtractionMode


def make_workbook(path: Path, sheets: int = 3) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for index in range(sheets):
        ws = wb.create_sheet(f"Sheet{index}")
        ws.append(["Name", "Count"])
        ws.append(["Item", 7])
        ws["A2"].hyperlink = "https://example.com"
        ws["B2"].fill = PatternFill("solid", fgColor="FF123456")
        ws["D1"] = "merged"
        ws.merge_cells("D1:E1")
        ws["F1"] = "=SUM(B2)"
        ws.print_area = "A1:F3"
        ws.add_table(Table(displayName=f"Table{index}", ref="A1:B2"))
        edge = Side(style="thin")
        for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=2):
            for cell in row:
                cell.border = Border(top=edge, bottom=edge, left=edge, right=edge)
    wb.save(path)
    wb.close()


def inputs_for(path: Path, *, formulas: bool = False) -> ExtractionInputs:
    return pipeline.resolve_extraction_inputs(
        path,
        mode="light",
        include_cell_links=True,
        include_print_areas=True,
        include_auto_page_breaks=False,
        include_colors_map=True,
        include_default_background=False,
        ignore_colors=set(),
        include_formulas_map=formulas,
        include_merged_cells=True,
        include_merged_values_in_rows=False,
    )


def track_loads(monkeypatch: pytest.MonkeyPatch) -> tuple[list[bool], list[Mock]]:
    original = workbook.load_workbook
    variants: list[bool] = []
    closes: list[Mock] = []

    def load(path: Path, *, data_only: bool, read_only: bool) -> Any:  # noqa: ANN401
        assert read_only is False
        variants.append(data_only)
        wb = original(path, data_only=data_only, read_only=read_only)
        close = Mock(wraps=wb.close)
        wb.close = close
        closes.append(close)
        return wb

    monkeypatch.setattr(workbook, "load_workbook", load)
    return variants, closes


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
@pytest.mark.parametrize("formulas", [False, True])
@pytest.mark.parametrize("sheet_count", [1, 4])
def test_pipeline_loads_once_per_variant_and_closes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    suffix: str,
    formulas: bool,
    sheet_count: int,
) -> None:
    path = tmp_path / f"book{suffix}"
    make_workbook(path, sheet_count)
    variants, closes = track_loads(monkeypatch)
    result = pipeline.run_extraction_pipeline(inputs_for(path, formulas=formulas))
    assert variants == ([True, False] if formulas else [True])
    assert len(result.workbook.sheets) == sheet_count
    for sheet in result.workbook.sheets.values():
        assert "A1:B2" in sheet.table_candidates
        assert sheet.print_areas
        assert sheet.merged_cells
        assert sheet.colors_map
        assert bool(sheet.formulas_map) is formulas
    assert all(close.call_count == 1 for close in closes)


def test_session_is_lazy_closed_and_rejects_reuse(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    variants, _ = track_loads(monkeypatch)
    session = OpenpyxlExtractionSession(tmp_path / "absent.xlsx")
    with session:
        assert variants == []
    with pytest.raises(RuntimeError, match="closed"):
        session.workbook()
    with pytest.raises(RuntimeError, match="closed"):
        session.__enter__()


@pytest.mark.parametrize("same_file", [False, True])
def test_com_table_detection_reuses_only_matching_session(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, same_file: bool
) -> None:
    """Equal sheet names must not select tables from a different workbook."""
    session_path = tmp_path / "session.xlsx"
    make_workbook(session_path, 1)
    target_path = tmp_path / "target.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet0"
    ws.append(["Name", "Count", "Other"])
    ws.append(["Item", 7, 9])
    ws.add_table(Table(displayName="Target", ref="A1:C2"))
    wb.save(target_path)
    wb.close()
    monkeypatch.chdir(tmp_path)
    sheet_path = session_path if same_file else target_path
    sheet = SimpleNamespace(
        name="Sheet0", book=SimpleNamespace(fullname=sheet_path.name)
    )
    expected = cells.detect_tables_openpyxl(sheet_path, sheet.name, mode="light")
    with OpenpyxlExtractionSession(session_path) as session:
        shared = session.workbook()
        loader = Mock(side_effect=AssertionError("session must already be loaded"))
        monkeypatch.setattr(session, "workbook", loader)
        if same_file:
            loader.side_effect = None
            loader.return_value = shared
        assert (
            cells.detect_tables(sheet, mode="light", openpyxl_session=session)
            == expected
        )
        assert loader.call_count == int(same_file)
    assert expected == (["A1:B2"] if same_file else ["A1:C2"])


def test_com_table_detection_preserves_helper_override_with_session(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "absent.xlsx"
    sheet = SimpleNamespace(name="Sheet0", book=SimpleNamespace(fullname=str(path)))
    helper = Mock(return_value=["D1:E2"])
    monkeypatch.setattr(cells, "detect_tables_openpyxl", helper)
    with OpenpyxlExtractionSession(path) as session:
        assert cells.detect_tables(sheet, mode="light", openpyxl_session=session) == [
            "D1:E2"
        ]
    helper.assert_called_once_with(path, "Sheet0", mode="light")


def test_close_all_variants_on_pipeline_model_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    _, closes = track_loads(monkeypatch)
    monkeypatch.setattr(
        pipeline, "build_workbook_data", Mock(side_effect=RuntimeError("model error"))
    )
    with pytest.raises(RuntimeError, match="model error"):
        pipeline.run_extraction_pipeline(inputs_for(path, formulas=True))
    assert len(closes) == 2
    assert all(close.call_count == 1 for close in closes)


def test_formula_load_failure_releases_regular_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    _, closes = track_loads(monkeypatch)
    original = workbook.load_workbook

    def fail_formula(path: Path, *, data_only: bool, read_only: bool) -> Any:  # noqa: ANN401
        if not data_only:
            raise ValueError("formula parse error")
        return original(path, data_only=data_only, read_only=read_only)

    monkeypatch.setattr(workbook, "load_workbook", fail_formula)
    result = pipeline.run_extraction_pipeline(inputs_for(path, formulas=True))
    assert result.workbook.sheets["Sheet0"].formulas_map == {}
    assert len(closes) == 1 and closes[0].call_count == 1


def test_enter_failure_closes_successfully_opened_resources(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    _, closes = track_loads(monkeypatch)

    class FailingSession(OpenpyxlExtractionSession):
        def __enter__(self) -> OpenpyxlExtractionSession:
            self.workbook()
            raise ValueError("enter failure")

    monkeypatch.setattr(pipeline, "OpenpyxlExtractionSession", FailingSession)
    with pytest.raises(ValueError, match="enter failure"):
        pipeline.run_extraction_pipeline(inputs_for(path))
    assert closes[0].call_count == 1


def test_standalone_helpers_match_shared_backend(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    standalone = OpenpyxlBackend(path)
    with OpenpyxlExtractionSession(path) as session:
        backend = OpenpyxlBackend(path, session=session)
        assert backend.extract_cells(include_links=True) == standalone.extract_cells(
            include_links=True
        )
        assert backend.extract_print_areas() == standalone.extract_print_areas()
        assert backend.extract_merged_cells() == standalone.extract_merged_cells()
        assert backend.extract_formulas_map() == standalone.extract_formulas_map()
        assert backend.extract_colors_map(
            include_default_background=True, ignore_colors=set()
        ) == standalone.extract_colors_map(
            include_default_background=True, ignore_colors=set()
        )
        for name in session.workbook().sheetnames:
            assert backend.detect_tables(name) == standalone.detect_tables(name)
            actual = cells.load_border_maps_openpyxl_ws(session.workbook()[name])
            expected = cells.load_border_maps_xlsx(path, name)
            for left, right in zip(actual[:5], expected[:5], strict=True):
                assert (left == right).all()
            assert actual[5:] == expected[5:]


def test_default_background_merged_cell_fallback_parity(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Keep the preexisting merged-cell color error and warning contract."""
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    assert (
        OpenpyxlBackend(path).extract_colors_map(
            include_default_background=True,
            ignore_colors=set(),
        )
        is None
    )
    with OpenpyxlExtractionSession(path) as session:
        assert (
            OpenpyxlBackend(path, session=session).extract_colors_map(
                include_default_background=True,
                ignore_colors=set(),
            )
            is None
        )
    assert caplog.text.count("Color map extraction failed; skipping colors_map") == 2


def test_all_feature_model_matches_standalone_helpers(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    inputs = inputs_for(path, formulas=True)
    shared = pipeline.run_extraction_pipeline(inputs).workbook
    monkeypatch.setattr(openpyxl_backend, "_use_session", lambda *_: False)
    standalone = pipeline.run_extraction_pipeline(inputs).workbook
    assert shared.model_dump() == standalone.model_dump()


def test_prestep_exception_releases_loaded_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    _, closes = track_loads(monkeypatch)

    def fail_after_load(
        inputs: ExtractionInputs, artifacts: pipeline.ExtractionArtifacts
    ) -> None:
        assert artifacts.openpyxl_session is not None
        artifacts.openpyxl_session.workbook()
        raise RuntimeError("prestep failure")

    monkeypatch.setattr(pipeline, "build_pre_com_pipeline", lambda _: [fail_after_load])
    with pytest.raises(RuntimeError, match="prestep failure"):
        pipeline.run_extraction_pipeline(inputs_for(path))
    assert len(closes) == 1 and closes[0].call_count == 1


def test_fresh_extraction_observes_file_changes(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path, 1)
    first = extract_workbook(path, mode="light")
    wb = workbook.load_workbook(path)
    wb["Sheet0"]["A2"] = "Changed"
    wb.save(path)
    wb.close()
    second = extract_workbook(path, mode="light")
    assert first.sheets["Sheet0"].rows != second.sheets["Sheet0"].rows


def test_concurrent_extractions_have_independent_workbooks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    barrier = Barrier(2)
    original = workbook.load_workbook
    loaded: list[object] = []

    def load(path: Path, **kwargs: Any) -> Any:  # noqa: ANN401
        wb = original(path, **kwargs)
        loaded.append(wb)
        barrier.wait(timeout=10)
        return wb

    monkeypatch.setattr(workbook, "load_workbook", load)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [
            pool.submit(pipeline.run_extraction_pipeline, inputs_for(path))
            for _ in range(2)
        ]
        results = [future.result(timeout=20) for future in futures]
    assert len(loaded) == 2 and loaded[0] is not loaded[1]
    assert results[0].workbook == results[1].workbook


@pytest.mark.parametrize("module", [cells, openpyxl_backend])
def test_legacy_helper_override_is_honored(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, module: ModuleType
) -> None:
    path = tmp_path / "absent.xlsx"
    fake = Mock(return_value={"Override": []})
    monkeypatch.setattr(module, "extract_sheet_cells", fake)
    with OpenpyxlExtractionSession(path) as session:
        assert OpenpyxlBackend(path, session=session).extract_cells(
            include_links=False
        ) == {"Override": []}
    fake.assert_called_once_with(path)


@pytest.mark.parametrize("success", [True, False])
def test_mock_com_success_and_fallback_share_session(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, success: bool
) -> None:
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    variants, closes = track_loads(monkeypatch)
    monkeypatch.delenv("SKIP_COM_TESTS", raising=False)
    inputs = replace(inputs_for(path), mode="standard")
    monkeypatch.setattr(pipeline, "build_com_pipeline", lambda _: [])

    @contextmanager
    def com_context(_: Path) -> Iterator[Any]:
        if not success:
            raise RuntimeError("COM unavailable")
        yield SimpleNamespace(
            sheets={
                name: SimpleNamespace(
                    name=name, book=SimpleNamespace(fullname=str(path))
                )
                for name in ("Sheet0", "Sheet1", "Sheet2")
            }
        )

    monkeypatch.setattr(pipeline, "xlwings_workbook", com_context)
    result = pipeline.run_extraction_pipeline(inputs)
    assert result.state.com_succeeded is success
    assert variants == [True]
    assert closes[0].call_count == 1
    assert all(
        "A1:B2" in sheet.table_candidates for sheet in result.workbook.sheets.values()
    )


@pytest.mark.parametrize("mode", ["light", "standard", "verbose"])
def test_public_xls_extraction_preserves_cells_and_com_formula_boundary(
    monkeypatch: pytest.MonkeyPatch,
    mode: ExtractionMode,
) -> None:
    path = Path(__file__).resolve().parents[1] / "assets" / "sample.xls"
    monkeypatch.setenv("SKIP_COM_TESTS", "1")
    original = workbook.load_workbook
    variants: list[bool] = []

    def load(path: Path, *, data_only: bool, read_only: bool) -> Any:  # noqa: ANN401
        variants.append(data_only)
        return original(path, data_only=data_only, read_only=read_only)

    monkeypatch.setattr(workbook, "load_workbook", load)
    result = extract_workbook(path, mode=mode)
    assert result.sheets and any(sheet.rows for sheet in result.sheets.values())
    assert False not in variants  # BIFF formulas must remain COM-only.
