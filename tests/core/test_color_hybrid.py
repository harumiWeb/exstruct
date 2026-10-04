"""Behavioral tests for the saved-fill plus rendered-color extraction path."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
import shutil
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Color, GradientFill, PatternFill
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet
import pytest

from exstruct.core import cells, color_hybrid, pipeline
from exstruct.core.cell_types import SheetColorsMap
from exstruct.core.openpyxl_session import OpenpyxlExtractionSession

_NO_FILL = -4142


def _excel_color(rgb: str) -> int:
    """Convert RGB hex to Excel's integer color representation."""
    red, green, blue = (int(rgb[index : index + 2], 16) for index in (0, 2, 4))
    return red | (green << 8) | (blue << 16)


class _FakeSheetAPI:
    def __init__(self, sheet: _FakeSheet) -> None:
        self.sheet = sheet

    def Activate(self) -> None:  # noqa: N802 - mirrors Excel COM
        pass

    def Calculate(self) -> None:  # noqa: N802 - mirrors Excel COM
        pass

    def Cells(self, row: int, column: int) -> SimpleNamespace:  # noqa: N802
        self.sheet.display_calls.append((row, column))
        color = self.sheet.rendered_colors.get((row, column), _NO_FILL)
        return SimpleNamespace(
            DisplayFormat=SimpleNamespace(
                Interior=_FakeInterior(self.sheet, (row, column), color)
            )
        )


class _FakeInterior:
    def __init__(
        self,
        sheet: _FakeSheet,
        coord: tuple[int, int],
        color: int | None,
    ) -> None:
        self.sheet = sheet
        self.coord = coord
        self.color = color

    @property
    def Color(self) -> int | None:  # noqa: N802 - mirrors Excel COM
        remaining = self.sheet.failures_remaining.get(self.coord, 0)
        if remaining:
            self.sheet.failures_remaining[self.coord] = remaining - 1
            raise RuntimeError("transient DisplayFormat failure")
        return self.color


class _FakeSheet:
    def __init__(
        self,
        name: str,
        bounds: tuple[int, int, int, int],
        rendered_colors: dict[tuple[int, int], int | None],
        failures_remaining: dict[tuple[int, int], int] | None = None,
    ) -> None:
        row, column, last_row, last_column = bounds
        self.name = name
        self.used_range = SimpleNamespace(
            row=row,
            column=column,
            last_cell=SimpleNamespace(row=last_row, column=last_column),
        )
        self.rendered_colors = rendered_colors
        self.failures_remaining = failures_remaining or {}
        self.display_calls: list[tuple[int, int]] = []
        self.api = _FakeSheetAPI(self)


class _FakeWorkbook:
    def __init__(
        self, fullname: Path, sheets: list[_FakeSheet], *, saved: bool = True
    ) -> None:
        self.fullname = str(fullname)
        self.sheets = sheets
        self.api = SimpleNamespace(Saved=saved)
        self.app = SimpleNamespace(calculate=lambda: None)


def _save_workbook(path: Path, configure: Callable[[Workbook], None]) -> None:
    workbook = Workbook()
    try:
        configure(workbook)
        workbook.save(path)
    finally:
        workbook.close()


def _fake_workbook(
    path: Path,
    *,
    used_ranges: dict[str, tuple[int, int, int, int]] | None = None,
    rendered_overrides: dict[str, dict[tuple[int, int], int | None]] | None = None,
    display_failures: dict[str, dict[tuple[int, int], int]] | None = None,
    fullname: Path | None = None,
    saved: bool = True,
) -> _FakeWorkbook:
    """Build COM-shaped sheets from a real saved workbook and optional Excel state."""
    used_ranges = used_ranges or {}
    rendered_overrides = rendered_overrides or {}
    display_failures = display_failures or {}
    fake_sheets = []
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        for worksheet in workbook.worksheets:
            min_column, min_row, max_column, max_row = range_boundaries(
                worksheet.calculate_dimension()
            )
            bounds = used_ranges.get(
                worksheet.title, (min_row, min_column, max_row, max_column)
            )
            rendered: dict[tuple[int, int], int | None] = {}
            row, column, last_row, last_column = bounds
            for row_index in range(row, last_row + 1):
                for column_index in range(column, last_column + 1):
                    fill = worksheet.cell(row_index, column_index).fill
                    if (
                        getattr(fill, "patternType", None) == "solid"
                        and fill.fgColor.type == "rgb"
                    ):
                        rgb = fill.fgColor.rgb[-6:]
                        rendered[row_index, column_index] = _excel_color(rgb)
                    else:
                        rendered[row_index, column_index] = _NO_FILL
            rendered.update(rendered_overrides.get(worksheet.title, {}))
            fake_sheets.append(
                _FakeSheet(
                    worksheet.title,
                    bounds,
                    rendered,
                    display_failures.get(worksheet.title),
                )
            )
    finally:
        workbook.close()
    return _FakeWorkbook(fullname or path, fake_sheets, saved=saved)


def _legacy_result(
    workbook: _FakeWorkbook,
    *,
    include_default: bool,
    ignore_colors: set[str] | None,
) -> dict[str, SheetColorsMap]:
    return {
        sheet.name: cells._extract_sheet_colors_com(
            sheet, include_default, ignore_colors
        )
        for sheet in workbook.sheets
    }


def _clear_display_calls(workbook: _FakeWorkbook) -> None:
    for sheet in workbook.sheets:
        sheet.display_calls.clear()


def _rewrite_sheet_xml(
    path: Path, sheet_name: str, transform: Callable[[bytes], bytes]
) -> None:
    with ZipFile(path) as archive:
        part = color_hybrid._worksheet_parts(archive)[sheet_name]
        members = [(name, archive.read(name)) for name in archive.namelist()]
    rewritten = []
    found = False
    for name, contents in members:
        if name == part:
            contents = transform(contents)
            found = True
        rewritten.append((name, contents))
    assert found
    temporary = path.with_name(f"{path.stem}.rewritten{path.suffix}")
    with ZipFile(temporary, "w", ZIP_DEFLATED) as archive:
        for name, contents in rewritten:
            archive.writestr(name, contents)
    temporary.replace(path)


def _add_expression_rule(worksheet: Worksheet, ref: str = "A1") -> None:
    worksheet.conditional_formatting.add(ref, FormulaRule(formula=["TRUE"]))


@pytest.mark.parametrize(
    ("include_default", "ignore_colors", "expected"),
    [
        (
            True,
            None,
            {
                "FFFFFF": [(1, 0), (1, 2), (2, 1)],
                "FF0000": [(1, 1)],
                "00FF00": [(2, 0)],
                "0000FF": [(2, 2)],
            },
        ),
        (
            False,
            None,
            {"FF0000": [(1, 1)], "00FF00": [(2, 0)], "0000FF": [(2, 2)]},
        ),
        (
            True,
            {"#ff0000", "ffffff"},
            {"00FF00": [(2, 0)], "0000FF": [(2, 2)]},
        ),
    ],
)
def test_static_fills_and_default_handling_match_legacy_com(
    tmp_path: Path,
    include_default: bool,
    ignore_colors: set[str] | None,
    expected: dict[str, list[tuple[int, int]]],
) -> None:
    path = tmp_path / "static.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"] = "plain"
        worksheet["B1"].fill = PatternFill("solid", fgColor="FFFF0000")
        worksheet["C1"].fill = PatternFill("solid", fgColor="FFFFFFFF")
        worksheet["A2"].fill = PatternFill("solid", fgColor="FF00FF00")
        worksheet["C2"].fill = PatternFill("solid", fgColor="FF0000FF")

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    legacy = _legacy_result(
        fake, include_default=include_default, ignore_colors=ignore_colors
    )
    _clear_display_calls(fake)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(
            fake, include_default, ignore_colors, session
        )

    assert actual.sheets == legacy
    assert actual.sheets["Sheet"].colors_map == expected
    assert fake.sheets[0].display_calls == []


def test_blank_cells_inside_com_used_range_keep_legacy_default_semantics(
    tmp_path: Path,
) -> None:
    path = tmp_path / "blank-extent.xlsx"

    def configure(workbook: Workbook) -> None:
        workbook.active["B2"] = "only saved value"

    _save_workbook(path, configure)
    # Excel's UsedRange can include blank cells outside openpyxl's saved dimension.
    fake = _fake_workbook(path, used_ranges={"Sheet": (2, 2, 3, 3)})
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets == legacy
    assert actual.sheets["Sheet"].colors_map == {
        "FFFFFF": [(2, 1), (2, 2), (3, 1), (3, 2)]
    }
    assert fake.sheets[0].display_calls == []


def test_conditional_candidates_are_deduplicated_clipped_and_row_major(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "conditional.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["B2"] = "top"
        worksheet["C3"] = "bottom"
        _add_expression_rule(worksheet, "A1:A3")
        _add_expression_rule(worksheet, "B2:B4")
        _add_expression_rule(worksheet, "C3:D8")
        _add_expression_rule(worksheet, "Z100:Z101")

    _save_workbook(path, configure)
    rendered = {
        "Sheet": {
            (2, 2): _excel_color("AA0000"),
            (3, 2): _excel_color("00AA00"),
            (3, 3): _excel_color("0000AA"),
        }
    }
    fake = _fake_workbook(path, rendered_overrides=rendered)
    legacy = _legacy_result(fake, include_default=False, ignore_colors=None)
    _clear_display_calls(fake)
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, False, None, session)

    assert actual.sheets == legacy
    assert fake.sheets[0].display_calls == [(2, 2), (3, 2), (3, 3)]
    assert actual.sheets["Sheet"].colors_map == {
        "AA0000": [(2, 1)],
        "00AA00": [(3, 1)],
        "0000AA": [(3, 2)],
    }
    log = next(
        record.getMessage()
        for record in caplog.records
        if "Color extraction sheet=Sheet" in record.getMessage()
    )
    assert "conditional_candidates=3" in log
    assert "rendered_candidates=3" in log
    assert "display_format_calls=3" in log
    assert "fallback=False" in log


@pytest.mark.parametrize(
    ("rendered_rgb", "ignore_colors", "expected"),
    [
        ("FFFFFF", None, {"FFFFFF": [(1, 0)]}),
        ("0000FF", {"#0000ff"}, {}),
    ],
)
def test_conditional_rendering_overrides_static_fill_before_ignore_filter(
    tmp_path: Path,
    rendered_rgb: str,
    ignore_colors: set[str] | None,
    expected: dict[str, list[tuple[int, int]]],
) -> None:
    path = tmp_path / "conditional-overrides-fill.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"].fill = PatternFill("solid", fgColor="FFFF0000")
        _add_expression_rule(worksheet)

    _save_workbook(path, configure)
    rendered = {"Sheet": {(1, 1): _excel_color(rendered_rgb)}}
    fake = _fake_workbook(path, rendered_overrides=rendered)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=ignore_colors)
    _clear_display_calls(fake)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, ignore_colors, session)

    assert actual.sheets == legacy
    assert actual.sheets["Sheet"].colors_map == expected
    assert fake.sheets[0].display_calls == [(1, 1)]


def test_theme_indexed_tinted_and_pattern_fills_use_rendered_candidates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "candidate-fills.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.25))
        worksheet["B1"].fill = PatternFill("solid", fgColor=Color(indexed=10))
        worksheet["C1"].fill = PatternFill(
            "solid",
            fgColor=Color(rgb="FF123456", tint=0.2),
            bgColor="FF000000",
        )
        worksheet["D1"].fill = PatternFill(
            patternType="darkGrid", fgColor="FFFF0000", bgColor="FF00FF00"
        )

    _save_workbook(path, configure)
    rendered = {
        "Sheet": {
            (1, 1): _excel_color("102030"),
            (1, 2): _excel_color("203040"),
            (1, 3): _excel_color("304050"),
            (1, 4): _excel_color("405060"),
        }
    }
    fake = _fake_workbook(path, rendered_overrides=rendered)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets == legacy
    assert fake.sheets[0].display_calls == [(1, 1), (1, 2), (1, 3), (1, 4)]
    assert actual.sheets["Sheet"].colors_map == {
        "102030": [(1, 0)],
        "203040": [(1, 1)],
        "304050": [(1, 2)],
        "405060": [(1, 3)],
    }


def test_transient_candidate_com_failure_discards_partial_scan_and_retries_legacy(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "transient-candidate-failure.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"] = "first candidate"
        worksheet["A2"] = "fails once"
        _add_expression_rule(worksheet, "A1:A2")

    _save_workbook(path, configure)
    rendered = {
        "Sheet": {
            (1, 1): _excel_color("AA1100"),
            (2, 1): _excel_color("0011AA"),
        }
    }
    legacy_fake = _fake_workbook(path, rendered_overrides=rendered)
    legacy = _legacy_result(legacy_fake, include_default=False, ignore_colors=None)
    hybrid_fake = _fake_workbook(
        path,
        rendered_overrides=rendered,
        display_failures={"Sheet": {(2, 1): 1}},
    )
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(hybrid_fake, False, None, session)

    assert actual.sheets == legacy
    assert actual.sheets["Sheet"].colors_map == {
        "AA1100": [(1, 0)],
        "0011AA": [(2, 0)],
    }
    assert hybrid_fake.sheets[0].display_calls == [
        (1, 1),
        (2, 1),
        (1, 1),
        (2, 1),
    ]
    assert cells._strict_display_format.get() is False
    assert any(
        "rendered candidate extraction failed: transient DisplayFormat failure"
        in record.getMessage()
        for record in caplog.records
    )
    log = next(
        record.getMessage()
        for record in caplog.records
        if "Color extraction sheet=Sheet" in record.getMessage()
    )
    assert "fallback=True" in log
    assert "display_format_calls=4" in log


@pytest.mark.parametrize(
    "raw_case",
    ["extLst", "unknown-rule", "foreign-namespace", "pivotTableParts"],
)
def test_untrusted_conditional_formatting_xml_uses_legacy_fallback(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    raw_case: str,
) -> None:
    path = tmp_path / f"raw-{raw_case}.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"] = "value"
        _add_expression_rule(worksheet)

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)

    def mutate(raw: bytes) -> bytes:
        if raw_case == "extLst":
            marker = b"</worksheet>"
            replacement = (
                b'<extLst><ext uri="urn:future"><x14:cf xmlns:x14='
                b'"http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'
                b"/></ext></extLst></worksheet>"
            )
        elif raw_case == "pivotTableParts":
            marker = b"</worksheet>"
            replacement = (
                b'<pivotTableParts count="1"><pivotTablePart r:id="rIdPivot" '
                b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
                b"/></pivotTableParts></worksheet>"
            )
        elif raw_case == "unknown-rule":
            marker = b'type="expression"'
            replacement = b'type="futureRule"'
        else:
            marker = b"<conditionalFormatting "
            replacement = b'<x:conditionalFormatting xmlns:x="urn:foreign" '
        assert marker in raw
        raw = raw.replace(marker, replacement, 1)
        if raw_case == "foreign-namespace":
            closing = b"</conditionalFormatting>"
            assert closing in raw
            raw = raw.replace(closing, b"</x:conditionalFormatting>", 1)
        return raw

    _rewrite_sheet_xml(path, "Sheet", mutate)
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets == legacy
    assert fake.sheets[0].display_calls == [(1, 1)]
    assert any(
        "Color extraction legacy fallback sheet=Sheet" in record.getMessage()
        for record in caplog.records
    )
    log = next(
        record.getMessage()
        for record in caplog.records
        if "Color extraction sheet=Sheet" in record.getMessage()
    )
    assert "fallback=True" in log


@pytest.mark.parametrize("fallback_case", ["unsaved", "session-mismatch", "non-xlsx"])
def test_unavailable_or_incompatible_saved_workbook_falls_back_to_legacy(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    fallback_case: str,
) -> None:
    path = tmp_path / "target.xlsx"

    def configure(workbook: Workbook) -> None:
        worksheet = workbook.active
        worksheet["A1"] = "start"
        worksheet["C2"].fill = PatternFill("solid", fgColor="FFABCDEF")

    _save_workbook(path, configure)
    fullname = path
    saved = fallback_case != "unsaved"
    session_path = path
    if fallback_case == "non-xlsx":
        fullname = tmp_path / "target.xls"
        shutil.copyfile(path, fullname)
    elif fallback_case == "session-mismatch":
        session_path = tmp_path / "other.xlsx"

        def configure_other(workbook: Workbook) -> None:
            workbook.active["A1"] = "other"

        _save_workbook(session_path, configure_other)

    fake = _fake_workbook(path, fullname=fullname, saved=saved)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    if fallback_case == "non-xlsx":
        actual = color_hybrid.extract_colors(fake, True, None, None)
    else:
        with OpenpyxlExtractionSession(session_path) as session:
            actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets == legacy
    assert fake.sheets[0].display_calls == [
        (1, 1),
        (1, 2),
        (1, 3),
        (2, 1),
        (2, 2),
        (2, 3),
    ]
    log = next(
        record.getMessage()
        for record in caplog.records
        if "Color extraction sheet=Sheet" in record.getMessage()
    )
    assert "fallback=True" in log


def test_fallback_is_isolated_to_the_sheet_with_unsupported_xml(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "sheet-isolation.xlsx"

    def configure(workbook: Workbook) -> None:
        broken = workbook.active
        broken.title = "Broken"
        broken["A1"] = "fallback"
        _add_expression_rule(broken)
        healthy = workbook.create_sheet("Healthy")
        healthy["A1"].fill = PatternFill("solid", fgColor="FF123ABC")

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)
    _rewrite_sheet_xml(
        path,
        "Broken",
        lambda raw: raw.replace(
            b"</worksheet>",
            b'<extLst><ext uri="urn:future"><x14:cf xmlns:x14='
            b'"http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'
            b"/></ext></extLst></worksheet>",
            1,
        ),
    )
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets["Broken"] == legacy["Broken"]
    assert actual.sheets["Healthy"] == legacy["Healthy"]
    assert fake.sheets[0].display_calls == [(1, 1)]
    assert fake.sheets[1].display_calls == []
    sheet_logs = {
        name: next(
            record.getMessage()
            for record in caplog.records
            if f"Color extraction sheet={name}" in record.getMessage()
        )
        for name in ("Broken", "Healthy")
    }
    assert "fallback=True" in sheet_logs["Broken"]
    assert "fallback=False" in sheet_logs["Healthy"]


def test_pipeline_com_color_step_reuses_shared_session_until_owner_closes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "shared-session.xlsx"

    def configure(workbook: Workbook) -> None:
        workbook.active["A1"] = "one"
        second = workbook.create_sheet("Second")
        second["B2"] = "two"

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)
    session = OpenpyxlExtractionSession(path)
    real_workbook = session.workbook
    workbook_calls: list[bool] = []
    loaded_workbook = None
    close_calls: list[int] = []

    def tracked_workbook(*, data_only: bool = True) -> object:
        nonlocal loaded_workbook
        workbook = real_workbook(data_only=data_only)
        workbook_calls.append(data_only)
        if loaded_workbook is None:
            loaded_workbook = workbook
            original_close = workbook.close

            def close() -> None:
                close_calls.append(1)
                original_close()

            workbook.close = close
        return workbook

    monkeypatch.setattr(session, "workbook", tracked_workbook)
    inputs = pipeline.resolve_extraction_inputs(
        path,
        mode="standard",
        include_cell_links=False,
        include_print_areas=False,
        include_auto_page_breaks=False,
        include_colors_map=True,
        include_default_background=True,
        ignore_colors=set(),
        include_formulas_map=False,
        include_merged_cells=False,
        include_merged_values_in_rows=False,
    )
    artifacts = pipeline.ExtractionArtifacts(openpyxl_session=session)
    backend_type = pipeline._backend_type("ComBackend")
    original_init = backend_type.__init__
    backend_sessions: list[OpenpyxlExtractionSession | None] = []

    def tracked_init(
        backend: object,
        workbook: object,
        session: OpenpyxlExtractionSession | None = None,
    ) -> None:
        backend_sessions.append(session)
        original_init(backend, workbook, session=session)

    monkeypatch.setattr(backend_type, "__init__", tracked_init)
    with session:
        pipeline.step_extract_colors_map_com(inputs, artifacts, fake)
        assert backend_sessions == [session]
        assert workbook_calls == [True, True]
        assert loaded_workbook is real_workbook()
        assert artifacts.colors_map_data is not None
        assert artifacts.colors_map_data.sheets == legacy
        assert all(sheet.display_calls == [] for sheet in fake.sheets)
        assert close_calls == []
        assert not session._closed

    assert session._closed
    assert close_calls == [1]


def test_custom_normal_fill_uses_rendered_legacy_fallback(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "custom-normal.xlsx"

    def configure(workbook: Workbook) -> None:
        workbook._named_styles["Normal"].fill = PatternFill("solid", fgColor="FF334455")
        workbook.active["A1"] = "inherits Normal"
        workbook.active["B2"] = "extends COM range"

    _save_workbook(path, configure)
    normal_color = _excel_color("334455")
    fake = _fake_workbook(
        path,
        used_ranges={"Sheet": (1, 1, 2, 2)},
        rendered_overrides={
            "Sheet": {
                (1, 1): normal_color,
                (1, 2): normal_color,
                (2, 1): normal_color,
                (2, 2): normal_color,
            }
        },
    )
    legacy = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)
    caplog.set_level("DEBUG", logger=color_hybrid.__name__)

    with OpenpyxlExtractionSession(path) as session:
        actual = color_hybrid.extract_colors(fake, True, None, session)

    assert actual.sheets == legacy
    assert actual.sheets["Sheet"].colors_map == {
        "334455": [(1, 0), (1, 1), (2, 0), (2, 1)]
    }
    assert fake.sheets[0].display_calls == [
        (1, 1),
        (1, 2),
        (2, 1),
        (2, 2),
    ]
    assert any(
        "Color extraction legacy fallback sheet=Sheet" in record.getMessage()
        and "custom Normal fill" in record.getMessage()
        for record in caplog.records
    )


def test_standalone_extraction_closes_the_session_it_creates(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "owned-session.xlsx"

    def configure(workbook: Workbook) -> None:
        workbook.active["A1"] = "value"

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    sessions: list[OpenpyxlExtractionSession] = []
    base_session = color_hybrid.OpenpyxlExtractionSession

    class TrackedSession(base_session):
        def __init__(self, file_path: Path) -> None:
            super().__init__(file_path)
            sessions.append(self)

    monkeypatch.setattr(color_hybrid, "OpenpyxlExtractionSession", TrackedSession)
    color_hybrid.extract_colors(fake, True, None, None)

    assert len(sessions) == 1
    assert sessions[0]._closed
    assert sessions[0]._workbooks == {}


@pytest.mark.parametrize("kind", ["merge", "table", "row", "column"])
def test_ambiguous_sheet_styles_preserve_legacy_scan(
    tmp_path: Path,
    kind: str,
) -> None:
    path = tmp_path / f"ambiguous-{kind}.xlsx"

    def configure(workbook: Workbook) -> None:
        ws = workbook.active
        ws.append(["Heading", "Value"])
        ws.append(["A", 1])
        fill = PatternFill("solid", fgColor="FF112233")
        if kind == "merge":
            ws.merge_cells("A1:B1")
        elif kind == "table":
            ws.add_table(Table(displayName="Data", ref="A1:B2"))
        elif kind == "row":
            ws.row_dimensions[1].fill = fill
        else:
            ws.column_dimensions["A"].fill = fill

    _save_workbook(path, configure)
    fake = _fake_workbook(
        path, rendered_overrides={"Sheet": {(1, 1): _excel_color("112233")}}
    )
    expected = _legacy_result(fake, include_default=True, ignore_colors=None)
    _clear_display_calls(fake)
    actual = color_hybrid.extract_colors(fake, True, None, None)
    assert actual.sheets == expected
    assert len(fake.sheets[0].display_calls) == 4


@pytest.mark.parametrize("kind", ["automatic", "gradient"])
def test_special_fills_require_excel_rendering(tmp_path: Path, kind: str) -> None:
    path = tmp_path / f"special-{kind}.xlsx"

    def configure(workbook: Workbook) -> None:
        workbook.active["A1"].fill = (
            PatternFill("solid", fgColor=Color(auto=True))
            if kind == "automatic"
            else GradientFill(stop=("FF112233", "FF445566"))
        )

    _save_workbook(path, configure)
    fake = _fake_workbook(
        path, rendered_overrides={"Sheet": {(1, 1): _excel_color("112233")}}
    )
    actual = color_hybrid.extract_colors(fake, False, None, None)
    assert actual.sheets["Sheet"].colors_map == {"112233": [(1, 0)]}
    assert fake.sheets[0].display_calls == [(1, 1)]


@pytest.mark.parametrize("problem", ["missing", "corrupt"])
def test_unreadable_saved_file_uses_live_legacy_colors(
    tmp_path: Path, problem: str
) -> None:
    path = tmp_path / "unreadable.xlsx"
    _save_workbook(path, lambda wb: wb.active.__setitem__("A1", "value"))
    fake = _fake_workbook(
        path, rendered_overrides={"Sheet": {(1, 1): _excel_color("112233")}}
    )
    if problem == "missing":
        path.unlink()
    else:
        path.write_bytes(b"invalid ZIP")
    actual = color_hybrid.extract_colors(fake, False, None, None)
    assert actual.sheets["Sheet"].colors_map == {"112233": [(1, 0)]}
    assert fake.sheets[0].display_calls == [(1, 1)]


@pytest.mark.parametrize("include_default", [False, True])
def test_sparse_static_colors_do_not_materialize_missing_cells(
    include_default: bool,
) -> None:
    """Default backgrounds must not enlarge the shared worksheet cell cache."""
    workbook = Workbook()
    ws = workbook.active
    ws["B2"].fill = PatternFill("solid", fgColor="FF123456")
    ws["C3"] = "unfilled"
    ws["D4"].fill = PatternFill("solid", fgColor=Color(theme=1))
    ws["E5"].fill = PatternFill("solid", fgColor="FFFFFFFF")
    ws["F6"].fill = PatternFill("solid", fgColor="FFABCDEF")
    stored = dict(ws._cells)
    candidates = {(1, 1)}
    try:
        actual = color_hybrid._static_colors(
            ws, (1, 1, 5, 5), candidates, include_default
        )
        assert ws._cells == stored
        assert candidates == {(1, 1), (4, 4)}
        assert actual[2, 2] == "123456"
        assert (6, 6) not in actual
        assert (1, 1) not in actual
        assert (4, 4) not in actual
        if include_default:
            assert len(actual) == 23
            assert actual[1, 2] == actual[3, 3] == actual[5, 5] == "FFFFFF"
        else:
            assert actual == {(2, 2): "123456"}
    finally:
        workbook.close()


def test_large_sparse_static_range_only_retains_colored_saved_cells() -> None:
    """A near-full-width UsedRange must not allocate its absent coordinates."""
    workbook = Workbook()
    ws = workbook.active
    ws["A1"] = "start"
    ws["XFD100000"].fill = PatternFill("solid", fgColor="FF123456")
    try:
        actual = color_hybrid._static_colors(ws, (1, 1, 100000, 16384), set(), False)
        assert actual == {(100000, 16384): "123456"}
        assert len(ws._cells) == 2
    finally:
        workbook.close()


@pytest.mark.parametrize("saved_state", [False, "unreadable"])
def test_sheet_preparation_invalidates_current_and_remaining_saved_fills(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    saved_state: bool | str,
) -> None:
    """Sheet event changes force live colors even if a later sheet appears saved."""
    path = tmp_path / "events.xlsm"

    def configure(workbook: Workbook) -> None:
        workbook.active.title = "Before"
        for name in ("Changed", "After"):
            workbook.create_sheet(name)
        for ws in workbook.worksheets:
            ws["A1"].fill = PatternFill("solid", fgColor="FF0000FF")

    _save_workbook(path, configure)
    fake = _fake_workbook(path)
    original_prepare = cells._prepare_sheet_for_display_format

    class UnreadableSaved:
        @property
        def Saved(self) -> bool:  # noqa: N802 - Excel COM property
            raise RuntimeError("Saved unavailable")

    def prepare(sheet: _FakeSheet) -> None:
        original_prepare(sheet)
        if sheet.name == "Changed":
            sheet.rendered_colors[1, 1] = _excel_color("FF0000")
            fake.sheets[2].rendered_colors[1, 1] = _excel_color("00FF00")
            fake.api = (
                SimpleNamespace(Saved=False)
                if saved_state is False
                else UnreadableSaved()
            )
        elif sheet.name == "After":
            fake.api = SimpleNamespace(Saved=True)

    monkeypatch.setattr(cells, "_prepare_sheet_for_display_format", prepare)
    actual = color_hybrid.extract_colors(fake, False, None, None)
    assert actual.sheets["Before"].colors_map == {"0000FF": [(1, 0)]}
    assert actual.sheets["Changed"].colors_map == {"FF0000": [(1, 0)]}
    assert actual.sheets["After"].colors_map == {"00FF00": [(1, 0)]}
    assert [sheet.display_calls for sheet in fake.sheets] == [[], [(1, 1)], [(1, 1)]]
