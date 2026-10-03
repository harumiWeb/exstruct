"""Cell contracts captured against the pandas reader before its replacement."""

from datetime import date, datetime, time, timedelta
import json
from pathlib import Path
import re
import subprocess
import sys
from unittest.mock import MagicMock
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink
import pytest
import xlrd

from exstruct.core import cells
from exstruct.models import CellRow


def make_parity_book(path: Path) -> None:
    """Cover reader conversions, missing tokens, cached formulas and row gaps."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Values"
    ws.append([123, 1.5, "00123", "1.50", True, False, "   ", " x "])
    ws.append([0, -5, 1e20, "1e2", "+001", "-.50", "1.", " NA "])
    ws.append(
        [
            date(2020, 1, 2),
            datetime(2020, 1, 2, 3, 4, 5, 123000),
            time(3, 4, 5, 123000),
            timedelta(days=2, hours=3),
        ]
    )
    tokens = [
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
    ]
    for col, token in enumerate(tokens, start=1):
        ws.cell(4, col, token).data_type = "s"
    for col, error in enumerate(
        ["#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"], start=1
    ):
        ws.cell(5, col, error)
    ws["D9"] = "sparse"
    ws["A10"] = "=1+1"
    ws["B10"] = "=TRUE()"
    ws["C10"] = '= "cached"'
    ws["D10"] = "=1/0"
    ws["E10"] = "=1+2"
    ws["A11"] = "visible"
    ws["B11"].hyperlink = "https://example.test/na"
    ws["B11"] = "NA"
    ws["A12"].hyperlink = "https://example.test/only"
    ws["A12"] = "NA"
    ws["C11"].hyperlink = Hyperlink(ref="C11", location="Values!A1")
    ws["C11"] = "internal"
    ws["A14"] = "first"
    ws.merge_cells("A14:B14")
    wb.create_sheet("Empty")
    wb.create_sheet("Last")["C3"] = "last"
    wb.save(path)
    wb.close()
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode().replace("<v />", "<v></v>")
    sheet = sheet.replace("<f>1+1</f><v></v>", "<f>1+1</f><v>2</v>")
    sheet = sheet.replace("<f>TRUE()</f><v></v>", "<f>TRUE()</f><v>1</v>")
    sheet = sheet.replace('r="B10"', 'r="B10" t="b"')
    sheet = sheet.replace('<f>= "cached"</f><v></v>', '<f>= "cached"</f><v>cached</v>')
    # The formula text is stored without its initial equals sign.
    sheet = sheet.replace('<f> "cached"</f><v></v>', '<f> "cached"</f><v>cached</v>')
    sheet = sheet.replace('r="C10"', 'r="C10" t="str"')
    sheet = sheet.replace("<f>1/0</f><v></v>", "<f>1/0</f><v>#DIV/0!</v>")
    sheet = sheet.replace('r="D10"', 'r="D10" t="e"')
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def dump_rows(data: dict[str, list[CellRow]]) -> dict[str, list[dict[str, object]]]:
    return {
        name: [row.model_dump(mode="json") for row in rows]
        for name, rows in data.items()
    }


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
def test_cells_match_frozen_old_reader(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"parity{suffix}"
    make_parity_book(path)
    expected = json.loads(
        (Path(__file__).parent / "cells_reader_expected.json").read_text(
            encoding="utf-8"
        )
    )
    assert dump_rows(cells.extract_sheet_cells(path)) == expected["xlsx"]
    assert dump_rows(cells.extract_sheet_cells_with_links(path)) == expected["links"]


def test_real_xls_matches_frozen_old_reader() -> None:
    path = Path(__file__).parents[1] / "assets" / "sample.xls"
    expected = json.loads(
        (Path(__file__).parent / "cells_reader_expected.json").read_text(
            encoding="utf-8"
        )
    )
    assert dump_rows(cells.extract_sheet_cells(path)) == expected["xls"]


def test_cells_match_live_pandas_oracle(tmp_path: Path) -> None:
    pd = pytest.importorskip("pandas")
    path = tmp_path / "oracle.xlsx"
    make_parity_book(path)
    expected: dict[str, list[CellRow]] = {}
    for name, frame in pd.read_excel(
        path, header=None, sheet_name=None, dtype=str
    ).items():
        rows = []
        for row_number, row in enumerate(
            frame.fillna("").itertuples(index=False, name=None), start=1
        ):
            values = {
                str(col): cells._coerce_numeric_preserve_format(str(value))
                for col, value in enumerate(row)
                if str(value).strip()
            }
            if values:
                rows.append(CellRow(r=row_number, c=values))
        expected[name] = rows
    assert cells.extract_sheet_cells(path) == expected


@pytest.mark.parametrize("datemode", [0, 1])
def test_xls_typed_cached_values_match_old_oracle(
    monkeypatch: pytest.MonkeyPatch, datemode: int
) -> None:
    """BIFF cell types include cached formulas; conversion ignores formula text."""
    pd = pytest.importorskip("pandas")
    types = [xlrd.XL_CELL_DATE] * 4 + [
        xlrd.XL_CELL_BOOLEAN,
        xlrd.XL_CELL_BOOLEAN,
        xlrd.XL_CELL_ERROR,
        xlrd.XL_CELL_NUMBER,
        xlrd.XL_CELL_NUMBER,
        xlrd.XL_CELL_NUMBER,
        xlrd.XL_CELL_TEXT,
        xlrd.XL_CELL_TEXT,
        xlrd.XL_CELL_TEXT,
        xlrd.XL_CELL_NUMBER,
        xlrd.XL_CELL_NUMBER,
    ]
    values = [
        0.5,
        1.0,
        59.0,
        60.0,
        1,
        0,
        7,
        2.0,
        1.25,
        1e20,
        "001",
        "NA",
        " NA ",
        float("inf"),
        float("nan"),
    ]
    row_cells = [
        xlrd.sheet.Cell(typ, value) for typ, value in zip(types, values, strict=True)
    ]
    sheet = MagicMock()
    sheet.name = "Cached"
    sheet.nrows = 3
    blank_cells = [xlrd.sheet.Cell(xlrd.XL_CELL_EMPTY, "") for _ in values]
    sheet.row.side_effect = lambda index: row_cells if index == 2 else blank_cells
    sheet.row_values.side_effect = lambda index: (
        values if index == 2 else [""] * len(values)
    )
    sheet.row_types.side_effect = lambda index: (
        types if index == 2 else [xlrd.XL_CELL_EMPTY] * len(values)
    )
    book = MagicMock(spec=xlrd.Book)
    book.datemode = datemode
    book.sheets.return_value = [sheet]
    book.sheet_names.return_value = ["Cached"]
    book.sheet_by_name.return_value = sheet
    monkeypatch.setattr(xlrd, "open_workbook", lambda *args, **kwargs: book)
    frame = pd.read_excel(
        book, engine="xlrd", sheet_name="Cached", header=None, dtype=str
    ).fillna("")
    expected: dict[str, int | float | str] = {
        str(column): cells._coerce_numeric_preserve_format(str(value))
        for column, value in enumerate(frame.iloc[2])
        if str(value).strip()
    }
    actual = cells.extract_sheet_cells(Path("cached.XLS"))
    assert actual == {"Cached": [CellRow(r=3, c=expected)]}
    assert actual["Cached"][0].c["0"] == "12:00:00"
    assert actual["Cached"][0].c["4"] == "True"
    assert actual["Cached"][0].c["5"] == "False"
    assert "6" not in actual["Cached"][0].c
    assert "11" not in actual["Cached"][0].c
    assert actual["Cached"][0].c["12"] == " NA "
    book.release_resources.assert_called()


def test_xls_release_on_read_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    book = MagicMock(spec=xlrd.Book)
    book.sheets.side_effect = RuntimeError("broken sheet")
    monkeypatch.setattr(xlrd, "open_workbook", lambda *args, **kwargs: book)
    with pytest.raises(RuntimeError, match="broken sheet"):
        cells.extract_sheet_cells(Path("broken.xls"))
    book.release_resources.assert_called_once()


def test_xls_date_overflow_keeps_old_serial() -> None:
    assert cells._xls_cell_value(1e20, xlrd.XL_CELL_DATE, 0) == 1e20


def test_ws_helper_matches_path_and_keeps_workbook_open(tmp_path: Path) -> None:
    path = tmp_path / "helper.xlsx"
    make_parity_book(path)
    with cells.openpyxl_workbook(path, data_only=True, read_only=False) as wb:
        rows = cells.extract_sheet_cells_openpyxl_ws(wb["Values"], include_links=True)
        assert rows == cells.extract_sheet_cells_with_links(path)["Values"]
        assert wb["Values"]["A10"].value == 2
        assert wb["Values"]["C10"].value == "cached"
        assert not any(row.r == 12 for row in rows)
        assert next(row for row in rows if row.r == 11).links == {
            "1": "https://example.test/na"
        }


def test_xlsx_works_without_pandas_and_keeps_xlrd_lazy(tmp_path: Path) -> None:
    path = tmp_path / "imports.xlsx"
    make_parity_book(path)
    code = (
        "import sys; sys.modules['pandas'] = None; from pathlib import Path; "
        "from exstruct.core.cells import extract_sheet_cells; "
        "extract_sheet_cells(Path(sys.argv[1])); "
        "assert sys.modules['pandas'] is None; assert 'xlrd' not in sys.modules"
    )
    subprocess.run([sys.executable, "-c", code, str(path)], check=True)


def test_streaming_reader_ignores_incorrect_declared_dimensions(tmp_path: Path) -> None:
    """The former pandas/openpyxl reader reset bounds before reading rows."""
    path = tmp_path / "dimensions.xlsx"
    wb = Workbook()
    wb.active["C3"] = "first"
    wb.active["D9"] = "last"
    wb.save(path)
    wb.close()
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = re.sub(
        rb'<dimension ref="[^"]+"',
        b'<dimension ref="A1:A1"',
        parts["xl/worksheets/sheet1.xml"],
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    assert cells.extract_sheet_cells(path) == {
        "Sheet": [CellRow(r=3, c={"2": "first"}), CellRow(r=9, c={"3": "last"})]
    }


def test_regular_reader_preserves_sparse_storage_and_coordinates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["C3"] = "first"
    sheet["AX1000"] = "last"
    sheet["AX1000"].hyperlink = "https://example.test/last"
    before = set(sheet._cells)
    assert cells.extract_sheet_cells_openpyxl_ws(sheet, include_links=True) == [
        CellRow(r=3, c={"2": "first"}),
        CellRow(r=1000, c={"49": "last"}, links={"49": "https://example.test/last"}),
    ]
    assert set(sheet._cells) == before
    workbook.close()
