"""Parity tests for OOXML table candidates without worksheet models."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side
from openpyxl.worksheet.table import Table
import pytest

from exstruct.core.backends.openpyxl_backend import OpenpyxlBackend
from exstruct.core.ooxml_session import OoxmlExtractionSession


def make_bordered_table_book(path: Path) -> None:
    """Create explicit and border-inferred tables plus inherited merge edges."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Name", "Count"])
    sheet.append(["one", 1])
    sheet.append(["two", 2])
    sheet.add_table(Table(displayName="ExplicitItems", ref="A1:B3"))

    edge = Side(style="thin")
    box = Border(top=edge, bottom=edge, left=edge, right=edge)
    for row in range(1, 5):
        for column in range(4, 7):
            cell = sheet.cell(row, column, f"value-{row}-{column}")
            cell.border = box

    sheet["H1"] = "merged"
    sheet["H1"].border = Border(top=edge, left=edge)
    sheet["I2"].border = Border(bottom=edge, right=edge)
    sheet.merge_cells("H1:I2")
    workbook.save(path)
    workbook.close()


def test_light_table_candidates_match_openpyxl_with_merge_border_inheritance(
    tmp_path: Path,
) -> None:
    path = tmp_path / "tables.xlsx"
    make_bordered_table_book(path)
    expected = OpenpyxlBackend(path).detect_tables("Data", mode="light")
    with OoxmlExtractionSession(path) as session:
        assert session.detect_tables("Data", mode="light") == expected
    assert "A1:B3" in expected
    assert "D1:F4" in expected


def test_border_only_candidates_match_openpyxl(tmp_path: Path) -> None:
    path = tmp_path / "border-only.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "BorderOnly"
    edge = Side(style="thin")
    border = Border(top=edge, bottom=edge, left=edge, right=edge)
    for row in range(1, 4):
        for column in range(3, 6):
            cell = sheet.cell(row, column, f"value-{row}-{column}")
            cell.border = border
    workbook.save(path)
    workbook.close()

    expected = OpenpyxlBackend(path).detect_tables("BorderOnly", mode="light")
    with OoxmlExtractionSession(path) as session:
        assert session.detect_tables("BorderOnly", mode="light") == expected
    assert "C1:E3" in expected


def test_detect_tables_reads_only_requested_sheets_table_parts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "many-tables.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "First"
    first.append(["Name", "Value"])
    first.append(["one", 1])
    first.add_table(Table(displayName="FirstItems", ref="A1:B2"))
    second = workbook.create_sheet("Second")
    second.append(["Name", "Value"])
    second.append(["two", 2])
    second.add_table(Table(displayName="SecondItems", ref="A1:B2"))
    workbook.save(path)
    workbook.close()

    with OoxmlExtractionSession(path) as session:
        archive = session.archive()
        original_read = archive.read
        table_reads: list[str] = []

        def track_read(name: str, *args: object, **kwargs: object) -> bytes:
            if name.startswith("xl/tables/"):
                table_reads.append(name)
            return original_read(name, *args, **kwargs)

        monkeypatch.setattr(archive, "read", track_read)
        assert session.detect_tables("First") == ["A1:B2"]
        assert table_reads == ["xl/tables/table1.xml"]
