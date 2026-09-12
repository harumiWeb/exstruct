"""Date-bearing cells must survive patch, inverse-op serialization, and undo."""

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from exstruct.edit import PatchOp, PatchRequest, patch_workbook
from exstruct.edit.types import coerce_patch_scalar


def _workbook_with_date(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Sheet1"
    sheet["A1"] = datetime(2025, 1, 1)
    workbook.save(path)
    workbook.close()


def test_patching_a_date_cell_captures_the_date_as_before(tmp_path: Path) -> None:
    """A date in the target cell used to fail PatchValue validation outright."""
    source = tmp_path / "book.xlsx"
    _workbook_with_date(source)

    result = patch_workbook(
        PatchRequest(
            xlsx_path=source,
            ops=[PatchOp(op="set_value", sheet="Sheet1", cell="A1", value="replaced")],
            backend="openpyxl",
            dry_run=True,
            return_inverse_ops=True,
        )
    )

    assert result.error is None
    assert result.patch_diff[0].before is not None
    assert result.patch_diff[0].before.value == datetime(2025, 1, 1)


def test_inverse_op_restores_a_real_date_through_json(tmp_path: Path) -> None:
    """The undo script is written to JSON, so the date hint must survive the trip."""
    source = tmp_path / "book.xlsx"
    patched = tmp_path / "patched.xlsx"
    undone = tmp_path / "undone.xlsx"
    _workbook_with_date(source)

    forward = patch_workbook(
        PatchRequest(
            xlsx_path=source,
            output_path=patched,
            ops=[PatchOp(op="set_value", sheet="Sheet1", cell="A1", value="replaced")],
            backend="openpyxl",
            return_inverse_ops=True,
        )
    )
    assert forward.error is None

    # Round-trip the inverse ops the way the CLI does: dump to JSON, read back.
    serialized = json.loads(json.dumps([op.model_dump(mode="json") for op in forward.inverse_ops]))
    assert forward.out_path is not None
    reverse = patch_workbook(
        PatchRequest(
            xlsx_path=Path(forward.out_path),
            output_path=undone,
            ops=[PatchOp.model_validate(op) for op in serialized],
            backend="openpyxl",
        )
    )
    assert reverse.error is None

    assert reverse.out_path is not None
    cell = load_workbook(reverse.out_path)["Sheet1"]["A1"]
    assert cell.value == datetime(2025, 1, 1)
    assert cell.data_type == "d", "undo downgraded the date cell to text"


def test_value_type_date_writes_a_date_not_a_string(tmp_path: Path) -> None:
    source = tmp_path / "book.xlsx"
    out = tmp_path / "out.xlsx"
    _workbook_with_date(source)

    result = patch_workbook(
        PatchRequest(
            xlsx_path=source,
            output_path=out,
            ops=[
                PatchOp(
                    op="set_value",
                    sheet="Sheet1",
                    cell="A1",
                    value="2026-12-25T00:00:00",
                    value_type="date",
                )
            ],
            backend="openpyxl",
        )
    )

    assert result.error is None
    assert result.out_path is not None
    assert load_workbook(result.out_path)["Sheet1"]["A1"].value == datetime(2026, 12, 25)


def test_value_type_auto_leaves_an_iso_string_alone() -> None:
    assert coerce_patch_scalar("2025-01-01", "auto") == "2025-01-01"


def test_value_type_date_rejects_non_iso_text() -> None:
    with pytest.raises(ValueError, match="not an ISO date/time"):
        coerce_patch_scalar("hello", "date")
