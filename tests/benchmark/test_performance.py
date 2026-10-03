"""Measurement contracts without machine-dependent timing thresholds."""

from contextlib import ExitStack
import json
from pathlib import Path
import subprocess
from typing import Any

from benchmark import performance
from benchmark.performance_fixtures import CATEGORIES, generate_fixtures
from benchmark.performance_profile import StageRecorder
import pytest


def test_fixtures_are_reproducible_and_refuse_overwrite(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    first = generate_fixtures(tmp_path / "first")
    second = generate_fixtures(tmp_path / "second")
    assert [path.stem for path in first] == list(CATEGORIES)
    for left, right in zip(first, second, strict=True):
        assert left.read_bytes() == right.read_bytes()
        workbook = load_workbook(left)
        sheets, rows, columns = CATEGORIES[left.stem]
        assert len(workbook.worksheets) == sheets
        assert workbook.worksheets[0].max_row == rows
        # small includes a chart outside the cell grid, which doesn't alter it.
        assert workbook.worksheets[0].max_column == columns
        workbook.close()
    with pytest.raises(FileExistsError):
        generate_fixtures(tmp_path / "first")


def test_recorder_restores_wrappers_after_exception() -> None:
    import openpyxl

    from exstruct.core.backends.openpyxl_backend import OpenpyxlBackend

    original = OpenpyxlBackend.extract_cells
    load = openpyxl.load_workbook
    with pytest.raises(RuntimeError), ExitStack() as stack:
        StageRecorder().install(stack)
        assert OpenpyxlBackend.extract_cells is not original
        assert openpyxl.load_workbook is not load
        raise RuntimeError("test")
    assert OpenpyxlBackend.extract_cells is original
    assert openpyxl.load_workbook is load


def test_timer_records_failure_without_swallowing_it() -> None:
    def fail() -> None:
        raise ValueError("expected")

    recorder = StageRecorder()
    with pytest.raises(ValueError, match="expected"):
        recorder.timed("cell_extraction", fail)()
    assert recorder.stages["cell_extraction"]["calls"] == 1


def test_fresh_worker_measures_light_and_standard_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import Workbook

    path = tmp_path / "tiny.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Count"])
    workbook.active.append(["Example", 7])
    workbook.save(path)
    workbook.close()
    monkeypatch.setenv("SKIP_COM_TESTS", "1")
    for mode in ("light", "standard"):
        _, output = performance.fresh_process(
            [
                "-m",
                "benchmark.performance",
                "--worker",
                "--input",
                str(path),
                "--mode",
                mode,
                "--repeats",
                "2",
            ],
            60,
        )
        result = json.loads(output)
        assert len(result["repeated"]) == 2
        assert result["modules_after_import"] == []
        assert isinstance(result["modules_after_extraction"], list)
        stages = result["profile"]["stages"]
        assert stages["cell_extraction"]["calls"] == 1
        assert stages["table_detection"]["calls"] == 1
        assert stages["model_construction"]["calls"] == 1
        assert stages["formula_extraction"]["calls"] == 0
        assert result["profile"]["openpyxl_workbook_opens"] >= 3
        assert result["profile"]["archive_opens"] >= 3
        state = result["profile"]["pipeline_state"]
        assert state["com_succeeded"] is False
        assert state["fallback_reason"] == (
            "skip_com_tests" if mode == "standard" else None
        )
        assert result["output_bytes"] > 0
        assert result["peak_rss_bytes"] is None or result["peak_rss_bytes"] > 0


@pytest.mark.parametrize(
    "arguments",
    [
        ["--repeats", "0"],
        ["--timeout", "nan"],
        ["--timeout", "inf"],
        ["--timeout", "0"],
        ["--input", "missing.xlsx"],
    ],
)
def test_invalid_arguments(arguments: list[str]) -> None:
    with pytest.raises(SystemExit) as error:
        performance.main(arguments)
    assert error.value.code == 2


def test_subprocess_errors_are_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "input.xlsx"
    path.write_bytes(b"placeholder")

    def fail(*args: Any, **kwargs: Any) -> dict[str, Any]:  # noqa: ANN401
        raise subprocess.TimeoutExpired("python", 1)

    monkeypatch.setattr(performance, "run_benchmark", fail)
    assert performance.main(["--input", str(path)]) == 1
    assert "Benchmark failed" in capsys.readouterr().err


def test_optional_profile_and_supervisor_report(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import Workbook

    path = tmp_path / "optional.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+2"
    workbook.save(path)
    workbook.close()
    monkeypatch.setenv("SKIP_COM_TESTS", "1")
    result = performance.run_benchmark(path, "light", 1, 60, profile_all_features=True)
    assert result["schema_version"] == 1
    assert result["mode"] == "light"
    assert len(result["input"]["sha256"]) == 64
    assert result["environment"]["SKIP_COM_TESTS"] == "1"
    assert result["profile"]["all_features"] is True
    for stage in ("formula_extraction", "color_extraction", "merged_cell_extraction"):
        assert result["profile"]["stages"][stage]["calls"] == 1
    assert result["startup"]["cold_extraction_process_ms"] > 0
    with pytest.raises(SystemExit):
        performance.main(["--input", str(path), "--output", str(path)])
