"""Issue 148 comparison; timings require a coordinator's serial slot.

Run correctness only: python -m benchmark.issue148_clustering --check-only
Run measurements: python -m benchmark.issue148_clustering --repeats 7 --output ...
All maps include grid-to-set construction in candidate timing. End-to-end
measurements include light extraction plus default JSON serialization, with
only the live clustering implementation overridden and restored afterwards.
"""

import argparse
from collections.abc import Callable
from datetime import UTC, datetime
from functools import partial
import hashlib
from importlib.metadata import PackageNotFoundError, version
import json
import os
from pathlib import Path
import platform
from statistics import median
import subprocess
import sys
from tempfile import TemporaryDirectory
from time import perf_counter
from typing import Any
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

import exstruct
from exstruct.core import cells

from .performance_fixtures import _normalize_archive

Cluster = Callable[[np.ndarray, int], list[tuple[int, int, int, int]]]
IMPLEMENTATIONS: dict[str, Cluster] = {
    "scipy": cells._detect_border_clusters_numpy,
    "python": cells._detect_border_clusters_python,
    "candidate_sparse": cells._detect_border_clusters_sparse,
}


def border_maps() -> dict[str, np.ndarray]:
    """Fixed occupied regions, including a large sparse grid and many labels."""
    small = np.zeros((40, 20), dtype=bool)
    small[2:8, 1:5] = True
    small[25:29, 12:17] = True
    large = np.zeros((2000, 200), dtype=bool)
    large[2:8, 1:5] = True
    large[1800:1805, 170:180] = True
    dense = np.ones((300, 80), dtype=bool)
    disconnected = np.zeros((300, 100), dtype=bool)
    for row in range(0, 300, 6):
        for col in range(0, 100, 5):
            disconnected[row : row + 2, col : col + 2] = True
    return {
        "sparse-small": small,
        "sparse-large": large,
        "dense-style-heavy": dense,
        "many-disconnected": disconnected,
    }


def write_fixture(path: Path, grid: np.ndarray) -> None:
    """Reproducible styled workbook preserving the complete input dimensions."""
    workbook = Workbook()
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    sheet = workbook.active
    assert sheet is not None
    border = Border(bottom=Side(style="thin"))
    rows, cols = np.nonzero(grid)
    for row, col in zip(rows.tolist(), cols.tolist(), strict=True):
        cell = sheet.cell(row + 1, col + 1, f"r{row}c{col}")
        cell.border = border
        if path.stem == "dense-style-heavy":
            cell.fill = PatternFill("solid", fgColor=f"{row * 313 % 0xFFFFFF:06X}")
            cell.font = Font(bold=row % 2 == 0, size=9 + col % 6)
            cell.number_format = "0.00"
    sheet.cell(grid.shape[0], grid.shape[1], "extent")
    workbook.save(path)
    workbook.close()
    _normalize_archive(path)


def observations(
    operation: Callable[[], Any], repeats: int, check_only: bool
) -> tuple[Any, dict[str, Any]]:
    """Warm correctness call followed by raw warm samples; no cold claim."""
    reference = operation()
    if check_only:
        return reference, {}
    samples: list[float] = []
    for _ in range(repeats):
        start = perf_counter()
        result = operation()
        samples.append((perf_counter() - start) * 1000)
        if result != reference:
            raise AssertionError("Output changed between measurement repetitions")
    return reference, {"warm_ms": samples, "median_ms": median(samples)}


def extraction(path: Path, implementation: Cluster) -> str:
    """Use existing python routing with a temporary live implementation patch."""
    with (
        patch.dict(os.environ, {"EXSTRUCT_BORDER_CLUSTER_BACKEND": "python"}),
        patch.object(cells, "_detect_border_clusters_python", implementation),
    ):
        return exstruct.serialize_workbook(
            exstruct.extract(path, mode="light"), fmt="json"
        )


def run_comparison(repeats: int, check_only: bool) -> dict[str, Any]:
    """Assert bbox and entire serialized output equality across all backends."""
    results: dict[str, Any] = {}
    with TemporaryDirectory(prefix="issue148-") as temporary:
        for category, grid in border_maps().items():
            path = Path(temporary) / f"{category}.xlsx"
            write_fixture(path, grid)
            case: dict[str, Any] = {
                "shape": list(grid.shape),
                "occupied": int(np.count_nonzero(grid)),
                "grid_bytes": grid.nbytes,
                "grid_sha256": hashlib.sha256(grid.tobytes()).hexdigest(),
                "fixture_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "min_size": 4,
                "cluster": {},
                "light_extraction_serialization": {},
            }
            expected_boxes = None
            expected_output = None
            for name, implementation in IMPLEMENTATIONS.items():
                boxes, timings = observations(
                    partial(implementation, grid, 4), repeats, check_only
                )
                if expected_boxes is None:
                    expected_boxes = boxes
                assert boxes == expected_boxes, (category, name, "bbox/order mismatch")
                case["cluster"][name] = timings
                serialized, timings = observations(
                    partial(extraction, path, implementation), repeats, check_only
                )
                if expected_output is None:
                    expected_output = serialized
                assert serialized == expected_output, (
                    category,
                    name,
                    "extraction mismatch",
                )
                case["light_extraction_serialization"][name] = timings
            assert expected_boxes is not None and expected_output is not None
            case["components"] = len(expected_boxes)
            case["output_sha256"] = hashlib.sha256(expected_output.encode()).hexdigest()
            case["output_bytes"] = len(expected_output.encode())
            case["bbox_order_equal"] = case["serialized_output_equal"] = True
            results[category] = case
    return results


def git_metadata(arguments: list[str]) -> str | None:
    try:
        result = subprocess.run(
            ["git", *arguments], capture_output=True, text=True, timeout=10, check=False
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    return result.stdout.strip() if result.returncode == 0 else None


def dependency_version(name: str) -> str | None:
    """Record optional absent dependencies without preventing comparison."""
    try:
        return version(name)
    except PackageNotFoundError:
        return None


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--repeats", type=int, default=7)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    status = git_metadata(["status", "--porcelain"])
    result = {
        "schema_version": 1,
        "timestamp_utc": datetime.now(UTC).isoformat(),
        "source_revision": git_metadata(["rev-parse", "HEAD"]),
        "source_dirty": None if status is None else bool(status),
        "python": sys.version,
        "platform": platform.platform(),
        "dependencies": {
            name: dependency_version(name)
            for name in ("numpy", "scipy", "pandas", "openpyxl")
        },
        "environment": {
            key: value
            for key, value in os.environ.items()
            if key.startswith("EXSTRUCT_")
        },
        "mode": "light",
        "warm_repeats": 0 if args.check_only else args.repeats,
        "check_only": args.check_only,
        "cases": run_comparison(args.repeats, args.check_only),
    }
    output = json.dumps(result, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output, encoding="utf-8")
    else:
        print(output, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
