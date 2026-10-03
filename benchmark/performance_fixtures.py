"""Fixed synthetic inputs for extraction performance comparisons."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

CATEGORIES = {
    "small": (1, 20, 8),
    "large": (1, 2000, 30),
    "many-sheet": (24, 30, 10),
    "sparse": (1, 1000, 50),
    "style-heavy": (1, 300, 20),
    "table-heavy": (1, 600, 12),
}


def _normalize_archive(path: Path) -> None:
    """Freeze ZIP timestamps as well as workbook properties for stable hashes."""
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            entry.date_time = (2000, 1, 1, 0, 0, 0)
            data = source.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                # openpyxl replaces modified at save time, so freeze the XML too.
                import re

                data = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\2",
                    data,
                )
            target.writestr(entry, data)
    path.write_bytes(output.getvalue())


def _write_fixture(path: Path, sheets: int, rows: int, columns: int) -> None:
    """Build and normalize one workbook in the staging directory."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table

    workbook = Workbook()
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    workbook.remove(workbook.active)
    for index in range(sheets):
        sheet = workbook.create_sheet(f"Sheet{index + 1:02d}")
        for column in range(1, columns + 1):
            sheet.cell(1, column, f"Column{column}")
        for row in range(2, rows + 1):
            for column in range(1, columns + 1):
                if path.stem == "sparse" and (row % 100 or column % 10):
                    continue
                cell = sheet.cell(row, column, row * columns + column)
                if path.stem == "style-heavy":
                    cell.fill = PatternFill(
                        "solid", fgColor=f"{row * 313 % 0xFFFFFF:06X}"
                    )
                    cell.font = Font(bold=row % 2 == 0, size=9 + column % 6)
                    cell.border = Border(bottom=Side(style="thin"))
                    cell.number_format = "0.00"
        sheet.print_area = f"A1:{sheet.cell(rows, columns).coordinate}"
        if path.stem == "table-heavy":
            for block in range(20):
                start = block * 30 + 1
                for column in range(1, columns + 1):
                    sheet.cell(start, column, f"Column{column}")
                sheet.add_table(
                    Table(
                        displayName=f"Table{block + 1}", ref=f"A{start}:L{start + 28}"
                    )
                )
        if path.stem == "small":
            sheet["H2"] = "=SUM(B2:G2)"
            sheet.merge_cells("A19:B19")
            chart = BarChart()
            chart.add_data(
                Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=10),
                titles_from_data=True,
            )
            sheet.add_chart(chart, "J2")
    workbook.save(path)
    workbook.close()
    _normalize_archive(path)


def generate_fixtures(directory: Path) -> list[Path]:
    """Stage complete fixtures before publishing; never overwrite existing files."""
    paths = [directory / f"{category}.xlsx" for category in CATEGORIES]
    if any(path.exists() for path in paths):
        raise FileExistsError("Fixture files already exist; use a new directory")
    directory.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=directory, prefix=".performance-") as temporary:
        staging = Path(temporary)
        for category, dimensions in CATEGORIES.items():
            _write_fixture(staging / f"{category}.xlsx", *dimensions)
        created: list[Path] = []
        try:
            for path in paths:
                with path.open("xb") as destination:
                    created.append(path)
                    destination.write((staging / path.name).read_bytes())
        except BaseException:
            for path in created:
                path.unlink()
            raise
    return paths
